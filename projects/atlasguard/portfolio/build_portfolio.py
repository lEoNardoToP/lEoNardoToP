from __future__ import annotations

import shutil, subprocess, textwrap, zipfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DInches, Pt as DPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent
OUT, FRAMES = ROOT / "generated", ROOT / ".frames"
VIDEOS, DOCS = OUT / "videos", OUT / "docs"
for d in (OUT, FRAMES, VIDEOS, DOCS): d.mkdir(parents=True, exist_ok=True)

PPTX = "AtlasGuard_Candidate_Portfolio_Canonical_SurveyMonkey_Sysdig.pptx"
PDF = "AtlasGuard_Candidate_Portfolio_Canonical_SurveyMonkey_Sysdig.pdf"
ZIP = "AtlasGuard_Candidate_Portfolio_Package.zip"

SLIDES = [
("AtlasGuard", ["Linux-native ML observability + explainable runtime security",
"Candidate engineering portfolio tailored to Canonical, SurveyMonkey and Sysdig",
"Leonardo Rosati - BSc Computer Science, University of Bologna (expected 2027)",
"Python • C/C++ • Java • Linux • Docker • FastAPI • ML evaluation • security tooling",
"6 tests passing • Python 3.11/3.12 CI • Docker build • PSI drift demo • explainable runtime rules"]),
("What I built", ["Observe - Linux host telemetry + inference telemetry enter a typed FastAPI contract.",
"Detect - PSI flags model-feature distribution shift; deterministic rules flag runtime anomalies.",
"Operate - Prometheus metrics, bounded state, failure-tolerant agent loop and reproducible demos.",
"Validation: 6 tests passing • PSI 7.83 after shift • 5 runtime findings • 2 Python versions in CI"]),
("Architecture", ["Linux agent -> typed telemetry contract -> FastAPI control plane",
"Inference service -> latency / feature signal -> drift detector",
"Runtime events -> explainable rule engine",
"Control plane -> bounded event store + Prometheus metrics"]),
("Engineering principles", ["Bounded - bounded deques instead of unbounded in-memory accumulation.",
"Resilient - network failures do not terminate collection.",
"Least privilege - non-root container, no added capabilities, read-only filesystem.",
"Explainable - stable rule IDs + human reasons; PSI is reproducible math.",
"Privacy-aware - metadata / numeric features, not raw prompts or user text.",
"Honest - eBPF, durable storage and distributed scale are roadmap items, not claimed features."]),
("Canonical fit", ["Target: Graduate Software Engineer - Open Source & Linux",
"Linux agent + systemd unit + container hardening",
"Typed Python modules, bounded state and failure-tolerant sender",
"CLI, quick-start, threat model and reproducible CI",
"Discussion: Ubuntu packaging, journald, profiling under load and upstream-friendly design."]),
("Canonical demo", ["24 seconds - Linux, validation, resilience and least privilege",
"Standalone video: videos/AtlasGuard_Canonical_Demo.mp4"]),
("SurveyMonkey fit", ["Target: Machine Learning Engineer - Italy / Remote",
"Model/service telemetry contract decoupled from model implementation",
"PSI baseline + rolling window surfaces statistical distribution shift",
"Prometheus latency + drift + system health metrics",
"Discussion: a model can degrade while the API remains healthy, so monitoring should be model-agnostic."]),
("Actual ML drift demonstration", ["Reference window vs shifted production window",
"PSI = 7.83", "Threshold = 0.20", "Result: DRIFTED = TRUE",
"The API remains healthy; the model-reliability signal changes."]),
("SurveyMonkey demo", ["24 seconds - telemetry, PSI and Prometheus signal",
"Standalone video: videos/AtlasGuard_SurveyMonkey_Demo.mp4"]),
("Sysdig fit", ["Target: Junior AI & Product Enablement Specialist - Italy / Remote",
"Docker/Python lab with repeatable normal + suspicious scenarios",
"Architecture + threat model + explicit non-goals",
"Linux/container runtime events with explainable findings",
"Discussion: build the lab, explain each finding, then swap synthetic events for Falco/eBPF collection."]),
("Sysdig-style hands-on lab", ["1 Start - control plane + metrics", "2 Observe - normal telemetry",
"3 Trigger - root shell / port 4444", "4 Explain - rule -> fields -> finding",
"5 Extend - swap in Falco/eBPF source"]),
("Sysdig demo", ["24 seconds - runtime event -> explainable finding",
"Standalone video: videos/AtlasGuard_Sysdig_Demo.mp4"]),
("What AtlasGuard does not claim", ["Not eBPF yet - runtime events are currently user-space inputs.",
"Not distributed storage - bounded in-memory state is deliberate for a reviewable prototype.",
"Not cloud-scale ML - this demonstrates platform mechanics, not AWS/Kafka/EKS/SageMaker production tenure.",
"The code tells the same story as the CV; roadmap items stay labeled as roadmap."]),
("What I would build next", ["Canonical: Debian/Ubuntu packaging, journald, profiling, upstream workflow, Ubuntu integration tests",
"SurveyMonkey: per-feature baselines, online evaluation, durable storage, OpenTelemetry, Kubernetes, SLO/load tests",
"Sysdig: Falco adapter, eBPF source, Kubernetes lab, workshop courseware, false-positive tuning exercise"]),
("Why I am showing you this", ["I am early-career, so I want the work to be inspectable.",
"AtlasGuard - ML observability + runtime security control plane",
"polyglot-alpha-beta-arena - algorithms + C++ autodiff + Java DQN + multi-language alpha-beta",
"Private R&D - HypoWeb scientific reasoning • IncidentLens defensive security analysis",
"The goal is not to look senior on paper. The goal is to show evidence of how I think as an engineer."]),
("Links, demos and sources", ["GitHub: github.com/lEoNardoToP/lEoNardoToP/tree/main/projects/atlasguard",
"LinkedIn: linkedin.com/in/leonardo-rosati-97b863273/",
"Email: leonardo.rosati@icloud.com", "Prepared 6 October 2026"])
]

BRIEFS = {
"Canonical": ("Graduate Software Engineer - Open Source & Linux",
["Linux host telemetry agent and systemd deployment.", "Typed Python control plane and resilient sender.",
"Bounded state and least-privilege container design.", "CI validates lint, tests and Docker on Python 3.11/3.12."],
["Ubuntu/Debian packaging", "journald integration", "CPU/memory profiling", "upstream contribution workflow"],
"Does not claim production-grade kernel visibility, distributed durability or eBPF collection."),
"SurveyMonkey": ("Machine Learning Engineer - Italy / Remote",
["Model-agnostic inference telemetry.", "PSI compares a reference distribution with a rolling window.",
"Drift is exported beside latency/system metrics.", "Raw prompts and user text are not stored."],
["per-feature/model baselines", "online evaluation hooks", "durable storage", "OpenTelemetry/Kubernetes", "SLO/load tests"],
"Demonstrates monitoring mechanics; does not claim AWS/Kafka/EKS/SageMaker production tenure."),
"Sysdig": ("Junior AI & Product Enablement Specialist - Italy / Remote",
["Repeatable Python/Docker lab scenarios.", "Architecture, threat model and explicit non-goals.",
"Runtime events map to stable rule IDs and human-readable findings.", "Direct path to a Falco/eBPF event source."],
["Falco adapter", "eBPF-backed source", "Kubernetes lab", "courseware", "false-positive tuning"],
"Not presented as a replacement for Sysdig or Falco; current runtime events are user-space inputs.")
}

def pilfont(size, bold=False):
    paths = [
      "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
      "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
    for p in paths:
        if Path(p).exists(): return ImageFont.truetype(p, size)
    return ImageFont.load_default()

def make_video(name, title, scenes):
    work = FRAMES / name
    if work.exists(): shutil.rmtree(work)
    work.mkdir(parents=True)
    frames = []
    for i, (heading, lines) in enumerate(scenes):
        img = Image.new("RGB", (1280,720), (10,16,28)); d = ImageDraw.Draw(img)
        d.rectangle((0,0,1280,10), fill=(73,189,255))
        d.text((70,60), title, fill=(73,189,255), font=pilfont(34,True))
        d.text((70,135), heading, fill=(245,248,252), font=pilfont(48,True))
        y=235
        for line in lines:
            first=True
            for part in textwrap.wrap(line, 62, break_long_words=False):
                d.text((90,y), ("• " if first else "  ")+part, fill=(205,216,232), font=pilfont(30))
                y += 48; first=False
        d.rounded_rectangle((70,645,1210,665),10,fill=(28,42,67))
        d.rounded_rectangle((70,645,70+int(1140*(i+1)/len(scenes)),665),10,fill=(74,222,128))
        d.text((70,682),"AtlasGuard • Leonardo Rosati",fill=(130,149,177),font=pilfont(18))
        p=work/f"scene_{i:02d}.png"; img.save(p); frames.append(p)
    concat=work/"concat.txt"
    with concat.open("w") as f:
        for p in frames: f.write(f"file '{p.as_posix()}'\nduration 6\n")
        f.write(f"file '{frames[-1].as_posix()}'\n")
    subprocess.run(["ffmpeg","-y","-loglevel","error","-f","concat","-safe","0","-i",str(concat),
      "-vf","fps=24,format=yuv420p","-c:v","libx264","-preset","veryfast","-movflags","+faststart",
      str(VIDEOS/f"{name}.mp4")],check=True)

def build_videos():
    make_video("AtlasGuard_Canonical_Demo","AtlasGuard × Canonical",[
      ("Linux-first design",["Host telemetry agent","systemd deployment","typed Python control plane"]),
      ("Resilience",["Controller failures do not terminate collection","bounded state","reproducible CLI demo"]),
      ("Least privilege",["non-root container","read-only filesystem","all Linux capabilities dropped"]),
      ("Engineering evidence",["Ruff + pytest","Python 3.11 / 3.12 matrix","Docker build in CI"])])
    make_video("AtlasGuard_SurveyMonkey_Demo","AtlasGuard × SurveyMonkey",[
      ("Production telemetry",["service + model identity","latency and confidence","numeric feature signal"]),
      ("Reference baseline",["100 observations create a baseline","API can stay healthy while model behavior changes"]),
      ("Distribution shift",["shifted production window","PSI = 7.83","threshold = 0.20"]),
      ("Operator signal",["drifted = true","Prometheus gauge exported","detector is replaceable"])])
    make_video("AtlasGuard_Sysdig_Demo","AtlasGuard × Sysdig",[
      ("Normal state",["Linux/container telemetry","no suspicious finding"]),
      ("Trigger",["privileged workload","uid=0 process","/bin/bash","egress port 4444"]),
      ("Explain",["stable rule identifier","human-readable reason","event fields map to finding"]),
      ("Extend",["swap synthetic source for Falco/eBPF","preserve control-plane contract","teach tuning"])])

def add_text(slide,text,x,y,w,h,size=18,bold=False,color=(245,248,252),align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; p.alignment=align
    r=p.runs[0]; r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*color)

def build_pptx():
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    for n,(title,items) in enumerate(SLIDES,1):
        s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(10,16,28)
        top=s.shapes.add_shape(1,0,0,prs.slide_width,Inches(.08)); top.fill.solid(); top.fill.fore_color.rgb=RGBColor(73,189,255); top.line.fill.background()
        add_text(s,title,.65,.55,11.8,.6,30 if n>1 else 44,True)
        y=1.45
        for i,b in enumerate(items):
            h=.68 if len(b)<110 else .9
            card=s.shapes.add_shape(5,Inches(.75),Inches(y),Inches(11.85),Inches(h))
            card.fill.solid(); card.fill.fore_color.rgb=RGBColor(20,31,51); card.line.color.rgb=RGBColor(37,55,84)
            add_text(s,b,1.0,y+.12,11.2,h-.16,18, i==0 and ("Target:" in b or n==1),(245,248,252))
            y += h+.16
            if y>6.65: break
        add_text(s,str(n),12.45,7.03,.35,.2,10,False,(171,184,204),PP_ALIGN.RIGHT)
    prs.save(OUT/PPTX)

def build_pdf():
    w,h=landscape(A4); c=Canvas(str(OUT/PDF),pagesize=(w,h))
    for n,(title,items) in enumerate(SLIDES,1):
        c.setFillColor(colors.HexColor("#0A101C")); c.rect(0,0,w,h,fill=1,stroke=0)
        c.setFillColor(colors.HexColor("#49BDFF")); c.rect(0,h-3*mm,w,3*mm,fill=1,stroke=0)
        c.setFillColor(colors.white); c.setFont("Helvetica-Bold",25); c.drawString(18*mm,h-25*mm,title)
        y=h-42*mm
        for b in items:
            c.setFillColor(colors.HexColor("#141F33")); c.roundRect(18*mm,y-16*mm,w-36*mm,18*mm,4*mm,fill=1,stroke=0)
            c.setFillColor(colors.HexColor("#E9EFF7")); c.setFont("Helvetica",11)
            ty=y-5*mm
            for line in textwrap.wrap(b,105)[:3]: c.drawString(23*mm,ty,line); ty-=5*mm
            y-=22*mm
            if y<20*mm: break
        c.setFillColor(colors.HexColor("#8194B1")); c.setFont("Helvetica",8); c.drawRightString(w-10*mm,8*mm,str(n)); c.showPage()
    c.save()

def build_docs():
    styles=getSampleStyleSheet()
    for company,(target,why,nxt,scope) in BRIEFS.items():
        doc=Document(); sec=doc.sections[0]; sec.top_margin=DInches(.65); sec.bottom_margin=DInches(.65); sec.left_margin=DInches(.75); sec.right_margin=DInches(.75)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(f"AtlasGuard × {company}"); r.bold=True; r.font.size=DPt(22)
        p=doc.add_paragraph(target); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for heading,vals in [("Why this project is relevant",why),("What I would extend next",nxt)]:
            p=doc.add_paragraph(); r=p.add_run(heading); r.bold=True; r.font.size=DPt(15)
            for x in vals: doc.add_paragraph(x,style="List Bullet")
        p=doc.add_paragraph(); r=p.add_run("Honest scope"); r.bold=True; r.font.size=DPt(15)
        doc.add_paragraph(scope); doc.add_paragraph("Author: Leonardo Rosati - BSc Computer Science, University of Bologna (in progress)")
        doc.save(DOCS/f"AtlasGuard_{company}_Technical_Brief.docx")
        title=ParagraphStyle("T",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=20,leading=24,alignment=TA_CENTER)
        h2=ParagraphStyle("H",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=13,leading=16,textColor=colors.HexColor("#166A99"))
        body=ParagraphStyle("B",parent=styles["BodyText"],fontName="Helvetica",fontSize=9.5,leading=13)
        pdf=SimpleDocTemplate(str(DOCS/f"AtlasGuard_{company}_Technical_Brief.pdf"),pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,topMargin=17*mm,bottomMargin=17*mm)
        story=[Paragraph(f"AtlasGuard × {company}",title),Spacer(1,6),Paragraph(target,body),Spacer(1,12),Paragraph("Why this project is relevant",h2)]
        for x in why: story += [Paragraph("• "+x,body),Spacer(1,3)]
        story += [Spacer(1,7),Paragraph("What I would extend next",h2)]
        for i,x in enumerate(nxt,1): story += [Paragraph(f"{i}. {x}",body),Spacer(1,3)]
        story += [Spacer(1,7),Paragraph("Honest scope",h2),Paragraph(scope,body),Spacer(1,12),Paragraph("Author: Leonardo Rosati - University of Bologna",body)]
        pdf.build(story)

def build_readme_zip():
    (OUT/"README_Portfolio_Package.md").write_text("""# AtlasGuard candidate portfolio package

Prepared for Leonardo Rosati - 6 October 2026.

- Main deck: PPTX + PDF
- Three tailored demo videos: Canonical, SurveyMonkey, Sysdig
- Three company-specific technical briefs: DOCX + PDF

AtlasGuard is a public portfolio prototype. Implemented functionality and roadmap items are deliberately separated.
""")
    z=OUT/ZIP
    with zipfile.ZipFile(z,"w",zipfile.ZIP_DEFLATED) as f:
        for p in sorted(OUT.rglob("*")):
            if p.is_file() and p!=z: f.write(p,p.relative_to(OUT))

def main():
    build_videos(); build_pptx(); build_pdf(); build_docs(); build_readme_zip()
    for p in sorted(OUT.rglob("*")):
        if p.is_file(): print(p.relative_to(OUT), p.stat().st_size)

if __name__ == "__main__": main()
